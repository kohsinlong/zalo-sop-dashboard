#!/usr/bin/env python3
"""
CZ weekly compliance review — build the DATA payload embedded in cz/index.html.

Source
  --tracker FILE.xlsx   The "NGT Cadence Compliance Tracker" workbook as
                        exported from Google Sheets (CZ_NGT_Tracker). Three
                        kinds of sheet are read:
                          CRM (live)     one row per customer, imported from the
                                         NGT Membership tracker: Zalo tag, name,
                                         DA, product, deposit / collection dates,
                                         CD anchor and the data flag
                          Cadence (CN)   what the DA should do at each milestone
                          W34, W35, …    one sheet per ISO week: the weekly
                                         summary CZ wrote at the top (sampled,
                                         pass, overall remark) and, per customer
                                         row, the two columns CZ fills in —
                                         "Missing cadence items" and
                                         "Qualitative feedback"
Output
  --write cz/index.html  Replace the `var DATA = …` line in place, or
  --json out.json        write the payload to a file.

Options
  --asof YYYY-MM-DD      Reporting date (default: today).
  --names full|none      Customer names are shown next to the Zalo tag so the
                         reviewer can find the conversation (default full).
                         Phone numbers and dates of birth are never embedded.

"Due this week" is NOT taken from the sheet: the page recomputes it from the
CD anchor and the product cadence with the same day offsets as the sheet's
formula, so any week can be chosen and the figures stay right as the CRM
changes. CZ's marks are kept as written; the verdict column (pass / partial /
fail / data issue) is read from column K of the week sheet when the page has
written one there, else derived from the wording of the two text columns.

Needs openpyxl (pip install openpyxl). Static page otherwise: no build step.
"""
import argparse, datetime, json, re, sys
from collections import Counter

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

# Day offsets from the CD anchor, exactly as the week sheets' "Due This Week"
# formula evaluates them. An N tag (glasses not collected yet) is only ever
# due D1 and D6-7; a Y tag follows its product's cadence.
OFFSETS = {
    "N":        [["D1", 1], ["D6-7", 7]],
    "MC":       [["CD", 0], ["CD+3-5", 5], ["CD+14", 14], ["CD+30", 30], ["CD+90", 90],
                 ["CD+180", 180], ["CD+270", 270], ["CD+365", 365]],
    "Kính":     [["CD", 0], ["CD+3-5", 5], ["CD+7", 7], ["CD+30", 30], ["CD+90", 90],
                 ["CD+180", 180], ["CD+270", 270], ["CD+365", 365]],
    "OK":       [["CD", 0], ["CDN", 0], ["CD+1", 1], ["CD+3-5", 5], ["CD+7", 7], ["CD+30", 30],
                 ["CD+90", 90], ["CD+180", 180], ["CD+270", 270], ["CD+365", 365]],
    "Atropine": [["CD", 0], ["CD+7", 7], ["CD+90", 90], ["CD+180", 180], ["CD+270", 270],
                 ["CD+365", 365]],
}

# The Cadence (CN) sheet is Chinese only; EN and VI are keyed by its text.
TRANSLATIONS = {
    "定金：发送介绍信息；更新预计配镜时间":
        ("Deposit: send the introduction; update the expected lens-ready date",
         "Đặt cọc: gửi thông tin giới thiệu; cập nhật thời gian dự kiến có kính"),
    "定金：发送介绍信息；更新时间":
        ("Deposit: send the introduction; update the timeline",
         "Đặt cọc: gửi thông tin giới thiệu; cập nhật thời gian"),
    "镜片到货后立即更新（<4小时）；告知取镜预约日":
        ("Update as soon as the lenses arrive (under 4 h); confirm the collection appointment",
         "Cập nhật ngay khi tròng kính về (dưới 4 giờ); báo ngày hẹn nhận kính"),
    "有药后更新；告知取药日":
        ("Update once the drops are in; confirm the pick-up date",
         "Cập nhật khi có thuốc; báo ngày nhận thuốc"),
    "发送使用说明；会员权益（客户回复后）；邀请转介绍":
        ("Send the usage guide; membership benefits (after the customer replies); invite referrals",
         "Gửi hướng dẫn sử dụng; quyền lợi hội viên (sau khi khách phản hồi); mời giới thiệu"),
    "取药：滴药指导；激活会员；发送权益（客户回复后）":
        ("Pick-up: instillation guidance; activate membership; send benefits (after the customer replies)",
         "Nhận thuốc: hướng dẫn nhỏ thuốc; kích hoạt hội viên; gửi quyền lợi (sau khi khách phản hồi)"),
    "取镜当天：检查首次佩戴Ortho-K；如不适及时协助":
        ("Collection night: check the first Ortho-K wear; help promptly if uncomfortable",
         "Đêm nhận kính: kiểm tra lần đeo Ortho-K đầu tiên; hỗ trợ ngay nếu khó chịu"),
    "检查情况；邀请转介绍":
        ("Check in; invite referrals", "Hỏi thăm; mời giới thiệu"),
    "3–5天后检查情况；邀请转介绍（客户回复后）":
        ("Check in after 3–5 days; invite referrals (after the customer replies)",
         "Hỏi thăm sau 3–5 ngày; mời giới thiệu (sau khi khách phản hồi)"),
    "检查情况；免费调整镜架；会员权益；邀请转介绍":
        ("Check in; free frame adjustment; membership benefits; invite referrals",
         "Hỏi thăm; chỉnh gọng miễn phí; quyền lợi hội viên; mời giới thiệu"),
    "电话检查依从性（≥每周6天）；提醒佩戴防眩光镜":
        ("Phone check on adherence (6+ days a week); remind to wear anti-glare glasses",
         "Gọi kiểm tra tuân thủ (từ 6 ngày/tuần); nhắc đeo kính chống chói"),
    "检查情况；提醒每天佩戴12小时；会员权益；邀请转介绍":
        ("Check in; remind to wear 12 h a day; membership benefits; invite referrals",
         "Hỏi thăm; nhắc đeo 12 giờ/ngày; quyền lợi hội viên; mời giới thiệu"),
    "检查情况；发送科普内容；会员权益；邀请转介绍":
        ("Check in; send educational content; membership benefits; invite referrals",
         "Hỏi thăm; gửi nội dung khoa học; quyền lợi hội viên; mời giới thiệu"),
    "检查情况；会员权益；邀请转介绍":
        ("Check in; membership benefits; invite referrals",
         "Hỏi thăm; quyền lợi hội viên; mời giới thiệu"),
    "邀请复查；会员权益；邀请转介绍":
        ("Invite for a re-examination; membership benefits; invite referrals",
         "Mời tái khám; quyền lợi hội viên; mời giới thiệu"),
    "邀请预约复查；会员权益；邀请转介绍":
        ("Invite to book a re-examination; membership benefits; invite referrals",
         "Mời đặt lịch tái khám; quyền lợi hội viên; mời giới thiệu"),
    "邀请预约复查；提醒每周佩戴>6次；会员权益；邀请转介绍":
        ("Invite to book a re-examination; remind to wear more than 6 nights a week; membership benefits; invite referrals",
         "Mời đặt lịch tái khám; nhắc đeo hơn 6 đêm/tuần; quyền lợi hội viên; mời giới thiệu"),
    "邀请复查；会员权益续购；终身免费调整镜架":
        ("Invite for a re-examination; membership renewal; lifetime free frame adjustment",
         "Mời tái khám; gia hạn hội viên; chỉnh gọng miễn phí trọn đời"),
    "复查1（3个月）；开具后续处方":
        ("Review 1 (3 months); issue the follow-on prescription",
         "Tái khám 1 (3 tháng); kê đơn tiếp theo"),
    "复查2（6个月）；咨询升级Ortho-K/MC":
        ("Review 2 (6 months); discuss upgrading to Ortho-K / MC",
         "Tái khám 2 (6 tháng); tư vấn nâng cấp Ortho-K / MC"),
    "复查3（9个月）；邀请转介绍":
        ("Review 3 (9 months); invite referrals",
         "Tái khám 3 (9 tháng); mời giới thiệu"),
    "复查4（12个月）；总结；咨询联合治疗方案":
        ("Review 4 (12 months); summary; discuss combined treatment",
         "Tái khám 4 (12 tháng); tổng kết; tư vấn phác đồ kết hợp"),
}

VERDICTS = ("pass", "partial", "fail", "data", "note")


# ---------------------------------------------------------------- helpers

def iso(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and 30000 < v < 60000:
        return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))).isoformat()
    s = str(v or "").strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return m.group(0)
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return "%s-%02d-%02d" % (m.group(3), int(m.group(2)), int(m.group(1)))
    return ""


def clean(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "").replace("\xa0", " ")).strip()


def first_line(v):
    return clean(str(v if v is not None else "").split("\n")[0])


def text(v):
    """Multi-line cell as typed, with stray quote marks and NBSPs tidied."""
    s = str(v if v is not None else "").replace("\xa0", " ").replace("\r", "")
    s = "\n".join(line.rstrip() for line in s.split("\n")).strip()
    return s.strip('"').strip()


def pad(row, n):
    """read_only sheets return ragged rows; pad to the header's width."""
    row = tuple(row)
    return row + (None,) * (n - len(row)) if len(row) < n else row


def col(hdr, needle):
    """First header starting with needle (whitespace-insensitive, case-insensitive)."""
    n = clean(needle).lower()
    for i, h in enumerate(hdr):
        if clean(h).lower().startswith(n):
            return i
    return -1


def product(v):
    """The sheet formula: SEARCH("Kính") anywhere → Kính; TRIM = OK / MC; Atropine."""
    s = clean(v)
    low = s.lower()
    if "kính" in low or "kinh" in low:
        return "Kính"
    if "atropine" in low or low.startswith("atr"):
        return "Atropine"
    if low == "ok":
        return "OK"
    if low == "mc":
        return "MC"
    return s


def tag_yn(tag):
    """Y / N after the product prefix: the sheet reads the letter just before
    the first digit of the compacted tag (MCY260720… → Y)."""
    s = re.sub(r"\s+", "", tag or "")
    m = re.search(r"\d", s)
    if not m or m.start() < 1:
        return ""
    for back in (1, 2):
        if m.start() - back >= 0 and s[m.start() - back].upper() in ("Y", "N"):
            return s[m.start() - back].upper()
    return ""


def tag_key(tag):
    """How a mark is matched to a customer: the compacted tag, lower-cased."""
    return re.sub(r"\s+", "", first_line(tag)).lower()


def monday_of(d):
    return (d - datetime.timedelta(days=d.weekday())).isoformat()


# ---------------------------------------------------------------- verdicts

RE_DATA = re.compile(r"wrong|incorrect|unclear|data flag|date not right|wrongly|labell?ed|无法查询|备注时间|备注名|标签错|日期不", re.I)
RE_PASS = re.compile(r"very good|well done|sent very good|^d6-?d?7:?\s*yes|called|很棒|完成|^sent$|^y$|^yes$", re.I)
RE_FAIL = re.compile(r"not sent|not send|:\s*no\b|：\s*n\b|:\s*n\b|^n[：:]|\bd30[：:]\s*n|no reply|未发送|未回访", re.I)
RE_HAS_Y = re.compile(r"(^|\n)\s*[\"“]?\s*(有|y)\s*[：:]", re.I)
RE_HAS_N = re.compile(r"(^|\n)\s*[\"“]?\s*(无|n)\s*[：:]", re.I)


def derive_verdict(missing, feedback):
    """pass / partial / fail / data / note from CZ's wording. Column K of the
    week sheet, when the page has written a verdict there, wins over this."""
    m = (missing or "").strip()
    both = (m + "\n" + (feedback or "")).strip()
    if not both:
        return ""
    if RE_DATA.search(m):
        return "data"
    has_y, has_n = bool(RE_HAS_Y.search(m)), bool(RE_HAS_N.search(m))
    if has_y and has_n:
        return "partial"
    if has_y:                       # a list of what was done, nothing flagged missing
        return "partial"
    if has_n:
        return "fail"
    if RE_PASS.search(m) and not RE_FAIL.search(m):
        return "pass"
    if RE_FAIL.search(m):
        sent = re.search(r"\b(sent|called)\b", re.sub(r"not\s+sen[dt]", "", m, flags=re.I), re.I)
        return "partial" if sent else "fail"
    return "note"


# ---------------------------------------------------------------- readers

def read_crm(wb, names):
    ws = wb["CRM (live)"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [clean(h) for h in rows[0]]
    c = {k: col(hdr, n) for k, n in {
        "no": "0", "friend": "Ngày kết bạn", "name": "Tên Khách Hàng", "prod": "Sản phẩm",
        "da": "Trợ lý", "tag": "Zalo", "status": "Trạng thái", "deposit": "Ngày đặt cọc",
        "coll": "Ngày nhận kính", "anchor": "Ngày neo", "flag": "Cờ cảnh báo", "kid2": "Tên bé thứ 2"}.items()}
    if c["tag"] < 0 or c["name"] < 0 or c["anchor"] < 0:
        sys.exit("CRM (live): need the Zalo, Tên Khách Hàng and Ngày neo columns")
    cases, seen = [], Counter()
    for r in rows[1:]:
        r = pad(r, len(hdr))
        tag = first_line(r[c["tag"]]) if c["tag"] >= 0 else ""
        name = clean(r[c["name"]])
        if not tag and not name:
            continue
        key = tag_key(tag)
        seen[key] += 1
        cases.append([
            int(r[c["no"]]) if c["no"] >= 0 and str(r[c["no"]] or "").strip().isdigit() else len(cases) + 1,
            tag,
            name if names == "full" else "",
            clean(r[c["da"]]) if c["da"] >= 0 else "",
            product(r[c["prod"]]) if c["prod"] >= 0 else "",
            iso(r[c["anchor"]]),
            iso(r[c["deposit"]]) if c["deposit"] >= 0 else "",
            iso(r[c["coll"]]) if c["coll"] >= 0 else "",
            iso(r[c["friend"]]) if c["friend"] >= 0 else "",
            tag_yn(tag),
            clean(r[c["flag"]]) if c["flag"] >= 0 else "",
            clean(r[c["kid2"]]) if c["kid2"] >= 0 and names == "full" else "",
        ])
    dup = sum(1 for k, n in seen.items() if k and n > 1)
    return cases, dup


def items(s):
    """One tick box per item: the sheet separates them with a full-width
    semicolon, the translations with a plain one."""
    return [x.strip() for x in re.split(r"[;；]", s) if x.strip()]


def read_cadence(wb):
    ws = wb["Cadence (CN)"]
    cad, missing = {}, []
    for r in list(ws.iter_rows(values_only=True))[1:]:
        r = pad(r, 4)
        if not r[0] or not r[1]:
            continue
        p, ms, zh = clean(r[0]), clean(r[1]), clean(r[2])
        en, vi = TRANSLATIONS.get(zh, ("", ""))
        if not en:
            missing.append(zh)
        elif not (len(items(zh)) == len(items(en)) == len(items(vi))):
            sys.exit("item count differs between languages for %s|%s: %r" % (p, ms, zh))
        cad.setdefault(p, {})[ms] = [items(zh), items(en) or items(zh), items(vi) or items(zh)]
    return cad, missing


RE_WEEK = re.compile(r"Week\s*(\d+)\s*[·•\-–—]*\s*(\d{4}-\d{2}-\d{2})\s*[–\-—]\s*(\d{4}-\d{2}-\d{2})")


def read_weeks(wb, cases):
    by_key = {}
    for c in cases:
        by_key.setdefault(tag_key(c[1]), c)
    weeks, carried, unmatched, seen_marks = [], 0, 0, {}
    for name in wb.sheetnames:
        if not re.fullmatch(r"W\d{1,2}", name):
            continue
        v = list(wb[name].iter_rows(values_only=True))
        m = RE_WEEK.search(clean(v[0][0]) if v and v[0] else "")
        v = [pad(row, 14) for row in v]
        if not m:
            print("  %s: no 'Week NN · from – to' title, skipped" % name)
            continue
        wk = {"k": m.group(2), "w": int(m.group(1)), "to": m.group(3),
              "sampled": None, "pass": None, "remark": "", "extra": "", "marks": []}
        for row in v[1:9]:
            for cell in pad(row, 14):
                s = text(cell)
                mm = re.match(r"^Sampled\s*[:：]\s*(\d+)", s, re.I)
                if mm:
                    wk["sampled"] = int(mm.group(1))
                mm = re.match(r"^Pass\s*[:：]\s*(\d+)", s, re.I)
                if mm:
                    wk["pass"] = int(mm.group(1))
        for row in v[1:9]:
            row = pad(row, 14)
            if re.search(r"Overall remark|综合评述", clean(row[0])):
                wk["remark"] = text(row[1]) if len(row) > 1 else ""
                extras = [text(x) for x in row[2:] if text(x) and text(x) != "-"]
                wk["extra"] = "\n".join(extras)
        hdr_i = next((i for i, row in enumerate(v) if row and re.search(r"Zalo Tag|Zalo 标签", clean(row[1]) if len(row) > 1 else "")), None)
        if hdr_i is None:
            weeks.append(wk)
            continue
        hdr = [clean(h) for h in v[hdr_i]]
        c_tag, c_miss, c_fb = 1, col(hdr, "缺失"), col(hdr, "质性")
        c_ver = c_fb + 1 if c_fb >= 0 else -1          # column K: the page writes its verdict here
        for row in v[hdr_i + 1:]:
            row = pad(row, len(hdr))
            if row[0] is None or not row[c_tag]:
                continue
            missing = text(row[c_miss]) if c_miss >= 0 else ""
            fb = text(row[c_fb]) if c_fb >= 0 else ""
            ver = clean(row[c_ver]).lower() if 0 <= c_ver < len(row) else ""
            if not missing and not fb and ver not in VERDICTS:
                continue
            key = tag_key(row[c_tag])
            sig = (key, missing, fb)
            if sig in seen_marks and seen_marks[sig] != wk["k"]:
                carried += 1                       # pasted forward from an earlier week, as written
                continue
            seen_marks.setdefault(sig, wk["k"])
            if key not in by_key:
                unmatched += 1
            wk["marks"].append([first_line(row[c_tag]), missing, fb,
                                ver if ver in VERDICTS else derive_verdict(missing, fb)])
        weeks.append(wk)
    weeks.sort(key=lambda w: w["k"])
    return weeks, carried, unmatched


# ---------------------------------------------------------------- main

def load_base(path):
    txt = open(path, encoding="utf-8").read()
    m = re.search(r"^(\s*var DATA = )(\{.*\});\s*$", txt, re.M)
    if not m:
        sys.exit("no `var DATA = {...};` line in " + path)
    return txt, m


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tracker", required=True)
    ap.add_argument("--write")
    ap.add_argument("--json")
    ap.add_argument("--asof", default=datetime.date.today().isoformat())
    ap.add_argument("--names", choices=["full", "none"], default="full")
    ap.add_argument("--reviewer", default="CZ")
    a = ap.parse_args()
    if not a.write and not a.json:
        ap.error("give --write or --json")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", a.asof):
        ap.error("--asof must be YYYY-MM-DD")

    wb = openpyxl.load_workbook(a.tracker, data_only=True, read_only=True)
    cases, dup = read_crm(wb, a.names)
    cad, untranslated = read_cadence(wb)
    weeks, carried, unmatched = read_weeks(wb, cases)

    payload = {"v": 1, "asof": a.asof, "generated": datetime.date.today().isoformat(),
               "reviewer": a.reviewer, "clinic": "NGT",
               "offsets": OFFSETS, "cad": cad, "weeks": weeks, "cases": cases}
    out = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))

    if a.json:
        open(a.json, "w", encoding="utf-8").write(out)
    if a.write:
        txt, m = load_base(a.write)
        open(a.write, "w", encoding="utf-8").write(txt[:m.start(2)] + out + txt[m.end(2):])

    das = Counter(c[3] for c in cases)
    prods = Counter(c[4] for c in cases)
    print("%d customers from CRM (live) (%d tags appear more than once) | DAs %s | products %s"
          % (len(cases), dup, dict(das), dict(prods)))
    marked = [w for w in weeks if w["marks"]]
    print("%d week sheets (%s to %s), %d with marks: %s"
          % (len(weeks), weeks[0]["k"] if weeks else "-", weeks[-1]["k"] if weeks else "-", len(marked),
             ", ".join("W%d=%d" % (w["w"], len(w["marks"])) for w in marked)))
    verdicts = Counter(mk[3] for w in weeks for mk in w["marks"])
    print("  verdicts derived: %s | marks carried forward unchanged and skipped: %d | marks whose tag is not in the CRM: %d"
          % (dict(verdicts), carried, unmatched))
    if untranslated:
        print("  cadence lines without EN/VI (shown in Chinese only): %d" % len(untranslated))
    print("  payload %d bytes" % len(out.encode("utf-8")))


if __name__ == "__main__":
    main()
