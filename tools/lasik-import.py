#!/usr/bin/env python3
"""
Lasik SOP dashboard — build the DATA payload embedded in lasik/index.html.

Sources
  --cte FILE.xlsx     A DA's customer-information workbook (the CTE format).
                      The surgery sheet (name ending in "_PT") gives one case
                      per row: surgery date (NGÀY PT), name (HỌ & TÊN), phone
                      (SĐT) and DA (NV TƯ VẤN). The exam sheet ("_Khám") says
                      whether the customer added Zalo; those who did not are
                      kept but not scored (verdict "nozalo").
  --audit FILE.xlsx   The D90 SOP QC workbook. Rows on its "D90 Audit" sheet
                      that match a case (phone + surgery date) supply the
                      touchpoint scores (10 / 5 / 0 per touchpoint).
  --base index.html   The current page. Cases of every clinic other than
                      --clinic are carried over unchanged (a v1 payload is
                      remapped onto the QC touchpoints first).

Modelling (until Salework is connected, like the optometry page)
  Cases without an audit row get modelled sent-flags: each touchpoint that
  has closed is "worked" with a probability calibrated on the existing
  modelled clinics (high at D1, tailing off by D60–90), scaled per DA. The
  per-DA scale is read off the carried-over data, so a DA's May looks like
  her June.
  --model-from YYYY-MM-DD  leave cases operated before this date "pending"
                           (unscored); default: model every unaudited case
  --no-model               keep every unaudited case pending
  --synth YYYY-MM=N        top the modelled clinics up to N cases in that
                           month with made-up ones, spread over their DAs
                           like the rest (MSG does about 1,000 Lasik
                           surgeries a month across the hospitals)
  --seed-ratings YYYY-MM:YYYY-MM  embed a QA response-quality rating (band 1-4 with
                           a justification, by Thao) for every DA for each month
                           in the range she had cases in. Later months stay
                           pending on the page. A comma list works too.

Output
  --write index.html  Replace the `var DATA = …` line in place, or
  --json out.json     write the payload to a file.

Options
  --clinic CTE              Clinic code for the --cte cases (default CTE).
  --asof YYYY-MM-DD         Reporting date. Surgeries after it are left out.
                            Defaults to the base page's as-of date.
  --phones last4|full|none  How much of each phone number to embed
                            (default last4: "…1234"). The page is public
                            once deployed, so full numbers are opt-in.

Needs openpyxl (pip install openpyxl). Static page otherwise: no build step.
"""
import argparse, datetime, hashlib, json, random, re, sys
from collections import Counter, defaultdict

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

# The QC guide's seven touchpoints: label, required items (key, content group),
# and the window in days after surgery during which the message is on time.
# Content groups feed the content KPIs: care follow-up, VIP benefits, referral.
TOUCH = [
    ["D1",     [["meds", "care"]],                          [1, 1]],
    ["D2–4",   [["checkin", "care"], ["vipintro", "vip"]],  [2, 4]],
    ["D6",     [["recall", "care"], ["referral", "ref"]],   [6, 6]],
    ["D15",    [["checkin", "care"]],                       [15, 15]],
    ["D25–30", [["recall1m", "care"], ["vipben", "vip"]],   [25, 30]],
    ["D30–60", [["checkin", "care"], ["eyecare", "care"]],  [30, 60]],
    ["D60–90", [["recall3m", "care"], ["referral", "ref"]], [60, 90]],
]
CP_TOTAL = sum(len(t[1]) for t in TOUCH)
# case row: [clinic, da, surgeryDate, verdict, scored, zaloTag, name, phone, ...flags]
CP_OFFSET = 8

# The v1 payload had 15 flags across seven older touchpoints. Modelled clinics
# are carried over by picking the closest older item for each QC item.
V1_MAP = [1, 2, 3, 4, 5, 6, 8, 9, 10, 11, 13, 14]

# Column headers on the audit sheet, matched by prefix (EN line of the header).
AUDIT_TP = ["D1", "D2–4", "D6", "D15", "D25–30", "D30–60", "D60–90"]

# Modelling: share of closed touchpoints that were worked at all, and the
# chance each item went out given that. Calibrated on the modelled clinics
# (D1 86%, D2–4 79%, D6 72%, D15 55%, D25–30 52%, D30–60 43%, D60–90 24%).
P_ANY = [0.86, 0.79, 0.72, 0.55, 0.52, 0.43, 0.24]
Q_ITEM = [[1.0], [0.99, 0.93], [0.99, 0.94], [1.0], [0.99, 0.80], [0.99, 0.80], [1.0, 0.83]]
NOZALO_SHARE = 0.06

FAMILY = ["Nguyễn", "Trần", "Lê", "Phạm", "Hoàng", "Huỳnh", "Phan", "Vũ", "Võ", "Đặng",
          "Bùi", "Đỗ", "Hồ", "Ngô", "Dương", "Lý", "Đinh", "Trịnh", "Lâm", "Mai"]
MIDDLE = ["Thị", "Văn", "Ngọc", "Minh", "Thanh", "Hoàng", "Gia", "Kim", "Bích", "Đức",
          "Quốc", "Xuân", "Hồng", "Thu", "Anh", "Bảo", "Khánh", "Phương", "Tuấn", "Hải"]
GIVEN = ["Anh", "Bảo", "Châu", "Dung", "Giang", "Hà", "Hân", "Hiếu", "Hoa", "Huy", "Khánh",
         "Lan", "Linh", "Long", "Mai", "Minh", "Nam", "Ngân", "Ngọc", "Nhi", "Phúc", "Phương",
         "Quân", "Quỳnh", "Sơn", "Thảo", "Thu", "Thư", "Trang", "Trâm", "Trung", "Tuấn", "Tú",
         "Uyên", "Vy", "Yến", "Đạt", "Hưng", "Kiệt", "Loan", "My", "Nga", "Oanh", "Tiên", "Vân"]

# The QA's justifications, one per band, varied per DA and month.
REMARKS = {
    1: ["Answered every message within the hour, in her own words rather than the script; reminded two customers about the 1-week check before they asked.",
        "Warm, specific replies; followed up the day after a customer mentioned dryness; VIP benefits explained naturally, not pasted.",
        "Proactive all month: chased the D25–30 reminders herself and handled a complaint about glare calmly and quickly.",
        "Customers replied to her check-ins with thanks; referral asks worded as a favour, never a push."],
    2: ["Timely and relevant replies; SOP content sent as required; one weekend message answered the next morning.",
        "Solid month. Answers matched the questions asked; VIP content occasionally pasted without personalising.",
        "Met expectations: a couple of generic replies, but nothing left unanswered and reminders on time.",
        "Good follow-through on the 1-month reminders; eye-care guidance could be more tailored to the customer."],
    3: ["Several replies took more than a day; two questions about eye drops went unanswered; scripts sent without adapting to the situation.",
        "Mostly copy-paste messages; a customer reporting discomfort waited until the next day for a reply.",
        "Follow-ups sent, but the tone was rushed and referral asks felt pushy; coaching agreed for next month.",
        "Reminders went out late in the window and check-ins were one-liners; customers stopped replying."],
    4: ["Customer messages unanswered for days; medication guidance sent to the wrong customer once; immediate coaching needed.",
        "Missed most conversations this month; replies late and off-topic; escalated to the clinic lead."],
}


def hash01(s):
    """Deterministic float in [0, 1) from a string."""
    return int(hashlib.sha256(s.encode("utf-8")).hexdigest()[:8], 16) / 2 ** 32


def rng(s):
    return random.Random(int(hashlib.sha256(s.encode("utf-8")).hexdigest()[:16], 16))


def norm_phone(v):
    """Digits only, leading zero restored (sheets store phones as numbers)."""
    if v is None:
        return ""
    s = str(v).strip()
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    d = re.sub(r"\D", "", s)
    if not d:
        return ""
    if d.startswith("84") and len(d) >= 11:
        d = "0" + d[2:]
    elif not d.startswith("0"):
        d = "0" + d
    return d


def mask_phone(p, mode):
    if not p:
        return ""
    if mode == "full":
        return p
    if mode == "none":
        return ""
    return "…" + p[-4:]


def iso(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        m = re.match(r"(\d{4})-(\d{2})-(\d{2})", v.strip())
        if m:
            return m.group(0)
        m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", v.strip())
        if m:
            return "%s-%02d-%02d" % (m.group(3), int(m.group(2)), int(m.group(1)))
    return ""


def clean_name(v):
    s = re.sub(r"\s+", " ", str(v or "")).strip()
    return s.title() if s.isupper() else s


def header_index(row, needle):
    """Column whose header equals needle; failing that, the first that starts
    with it. Exact first, because "Họ & tên KH giới thiệu" (the referrer)
    sits two columns before "HỌ & TÊN" (the customer)."""
    def norm(h):
        return re.sub(r"\s+", " ", str(h).replace("\n", " ")).strip().lower()
    n = needle.lower()
    for i, h in enumerate(row):
        if h and norm(h) == n:
            return i
    for i, h in enumerate(row):
        if h and norm(h).startswith(n):
            return i
    return -1


def days_between(a, b):
    return (datetime.date.fromisoformat(b) - datetime.date.fromisoformat(a)).days


# ---------------------------------------------------------------- sources

def read_cte(path, clinic, phones):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    pt = next((s for s in wb.sheetnames if s.endswith("_PT")), None)
    kh = next((s for s in wb.sheetnames if s.endswith("_Khám")), None)
    if not pt:
        sys.exit("no surgery sheet (name ending in _PT) in " + path)

    rows = list(wb[pt].iter_rows(values_only=True))
    hdr = rows[0]
    c_date = header_index(hdr, "NGÀY PT")
    c_name = header_index(hdr, "HỌ & TÊN")
    c_phone = header_index(hdr, "SĐT")
    c_da = header_index(hdr, "NV TƯ VẤN")
    if min(c_date, c_name, c_phone) < 0:
        sys.exit("surgery sheet is missing NGÀY PT / HỌ & TÊN / SĐT columns")

    zalo = {}
    if kh:
        krows = list(wb[kh].iter_rows(values_only=True))
        k_phone = header_index(krows[0], "SDT")
        if k_phone < 0:
            k_phone = header_index(krows[0], "SĐT")
        k_zalo = header_index(krows[0], "Zalo")
        if k_phone >= 0 and k_zalo >= 0:
            for r in krows[1:]:
                p = norm_phone(r[k_phone]) if k_phone < len(r) else ""
                z = str(r[k_zalo] or "").strip().lower() if k_zalo < len(r) else ""
                if p and z:
                    zalo[p] = "y" if z.startswith("c") else "n" if z.startswith("k") else ""

    das = Counter(str(r[c_da]).strip() for r in rows[1:] if c_da >= 0 and r[c_da])
    da_default = das.most_common(1)[0][0] if das else ""

    cases = []
    for r in rows[1:]:
        d = iso(r[c_date]) if c_date < len(r) else ""
        name = clean_name(r[c_name]) if c_name < len(r) else ""
        if not d or not name:
            continue
        phone = norm_phone(r[c_phone]) if c_phone < len(r) else ""
        da = str(r[c_da]).strip() if c_da >= 0 and c_da < len(r) and r[c_da] else da_default
        verdict = "nozalo" if zalo.get(phone) == "n" else "pending"
        cases.append({
            "clinic": clinic, "da": da, "date": d, "verdict": verdict,
            "name": name, "phone": phone, "show": mask_phone(phone, phones),
            "flags": [0] * CP_TOTAL,
        })
    return cases, {"sheet": pt, "rows": len(cases), "zaloKnown": len(zalo)}


def read_audit(path):
    """{(phone, surgeryDate): [7 touchpoint scores]} from the D90 Audit sheet."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    name = next((s for s in wb.sheetnames if s.lower().startswith("d90 audit")), None)
    if not name:
        return {}
    rows = list(wb[name].iter_rows(values_only=True))
    hdr_i = next((i for i, r in enumerate(rows[:10])
                  if any(h and str(h).startswith("Phone") for h in r)), None)
    if hdr_i is None:
        return {}
    hdr = rows[hdr_i]
    c_phone = header_index(hdr, "Phone")
    c_date = header_index(hdr, "Surgery")
    c_tp = [header_index(hdr, lbl) for lbl in AUDIT_TP]
    out = {}
    for r in rows[hdr_i + 1:]:
        p = norm_phone(r[c_phone]) if c_phone < len(r) else ""
        d = iso(r[c_date]) if c_date < len(r) else ""
        if not p or not d:
            continue
        scores = []
        for i in c_tp:
            v = r[i] if 0 <= i < len(r) else None
            try:
                scores.append(int(float(v)) if v not in (None, "") else None)
            except ValueError:
                scores.append(None)
        out[(p, d)] = scores
    return out


def apply_audit(cases, audit):
    n = 0
    for c in cases:
        s = audit.get((c["phone"], c["date"]))
        if s is None:
            continue
        flags = []
        for ti, t in enumerate(TOUCH):
            k = len(t[1])
            v = s[ti]
            if v is None or v <= 0:
                flags += [0] * k
            elif v >= 10:
                flags += [1] * k
            else:                       # 5 = one of two items: mark the first
                flags += [1] + [0] * (k - 1)
        c["flags"] = flags
        c["verdict"] = "audited"
        n += 1
    return n


def load_base(path):
    txt = open(path, encoding="utf-8").read()
    m = re.search(r"^(\s*var DATA = )(\{.*\});\s*$", txt, re.M)
    if not m:
        sys.exit("no `var DATA = {...};` line in " + path)
    return txt, m, json.loads(m.group(2))


def carry_over(base, clinic):
    """Cases of the other clinics, in the current row shape."""
    out = []
    v2 = base.get("v", 1) >= 2
    for c in base["cases"]:
        if c[0] == clinic:
            continue
        if v2:
            out.append(c)
        else:   # v1: [clinic, da, date, verdict, scored, tag, ...15 flags]
            flags = c[6:]
            out.append(c[:6] + ["", ""] + [flags[i] if i < len(flags) else 0 for i in V1_MAP])
    return out


# ---------------------------------------------------------------- modelling

def da_multipliers(rows, asof):
    """How much more or less than the average each DA works her closed
    touchpoints, read off the carried-over modelled data (clamped)."""
    worked, closed = Counter(), Counter()
    for c in rows:
        if c[4] != 1:
            continue
        age = days_between(c[2], asof)
        for ti, t in enumerate(TOUCH):
            if age <= t[2][1]:
                continue
            k = (c[0], c[1])
            closed[k] += 1
            if any(c[CP_OFFSET + sum(len(x[1]) for x in TOUCH[:ti]) + i] for i in range(len(t[1]))):
                worked[k] += 1
    tot_w, tot_c = sum(worked.values()), sum(closed.values())
    if not tot_c:
        return {}
    avg = tot_w / tot_c
    return {k: max(0.75, min(1.25, (worked[k] / closed[k]) / avg))
            for k in closed if closed[k] >= 20}


def model_flags(key, mul, date, asof):
    """Sent-flags for one case: closed touchpoints are worked with P_ANY
    scaled by the DA, each item then goes out with its own chance; an open
    window is left unsent (it shows as "Open now" on the page)."""
    r = rng("flags|" + key)
    age = days_between(date, asof)
    case_mul = 0.9 + 0.2 * r.random()
    flags = []
    for ti, t in enumerate(TOUCH):
        k = len(t[1])
        if age <= t[2][1] or r.random() >= P_ANY[ti] * mul * case_mul:
            flags += [0] * k
            continue
        bits = [1 if r.random() < Q_ITEM[ti][i] else 0 for i in range(k)]
        if not any(bits):
            bits[0] = 1
        flags += bits
    return flags


def verdict_for(flags, date, asof):
    age = days_between(date, asof)
    full = part = miss = 0
    off = 0
    for ti, t in enumerate(TOUCH):
        k = len(t[1])
        if age > t[2][1]:
            got = sum(flags[off:off + k])
            if got == k:
                full += 1
            elif got:
                part += 1
            else:
                miss += 1
        off += k
    if not (full or part or miss):
        return "tooearly"
    if not (part or miss):
        return "pass"
    return "none" if not (full or part) else "partial"


def model_cases(cases, da_mul, model_from, asof):
    """Modelled sent-flags for the --cte cases operated from model_from on."""
    n = 0
    for c in cases:
        if c["verdict"] != "pending" or c["date"] < model_from:
            continue
        mul = da_mul.get((c["clinic"], c["da"]), 1.0)
        c["flags"] = model_flags(c["clinic"] + "|" + c["da"] + "|" + c["date"] + "|" + c["phone"] + c["name"],
                                 mul, c["date"], asof)
        c["verdict"] = verdict_for(c["flags"], c["date"], asof)
        n += 1
    return n


def fake_tag(r):
    """A Zalo tag in the style of the modelled clinics' data."""
    name = " ".join([r.choice(FAMILY), r.choice(MIDDLE), r.choice(GIVEN)])
    style = r.random()
    if style < 0.6:
        return name + " 0" + "".join(str(r.randint(0, 9)) for _ in range(9))
    if style < 0.8:
        return r.choice(MIDDLE) + " " + r.choice(GIVEN)
    return r.choice(GIVEN) + " " + r.choice(GIVEN) + " - " + r.choice(["Femto", "Lasik", "Smile", "Phakic"])


def synth_cases(rows, month, n, da_mul, asof, clinic_skip):
    """Top the modelled clinics up to N cases in `month` with made-up ones,
    spread over their DAs in the same proportions as the carried-over data.
    Re-running with the same N adds nothing."""
    mix = Counter((c[0], c[1]) for c in rows if c[0] != clinic_skip)
    total = sum(mix.values())
    existing = sum(1 for c in rows if c[0] != clinic_skip and c[2].startswith(month))
    n = max(0, n - existing)
    if not total or not n:
        return []
    r = rng("synth|" + month + "|" + str(existing))
    quota = {k: round(n * v / total) for k, v in mix.items()}
    y, m = int(month[:4]), int(month[5:7])
    last = (datetime.date(y + (m == 12), (m % 12) + 1, 1) - datetime.timedelta(days=1)).day
    if month == asof[:7]:                      # the live month runs to the as-of date only
        last = int(asof[8:10])
    out, seen = [], set(c[5] for c in rows)
    for (cl, da), q in sorted(quota.items()):
        for i in range(q):
            date = "%s-%02d" % (month, r.randint(1, last))
            if date > asof:
                continue
            tag = fake_tag(r)
            while tag in seen:
                tag = fake_tag(r)
            seen.add(tag)
            if r.random() < NOZALO_SHARE:
                out.append([cl, da, date, "nozalo", 0, tag, "", ""] + [0] * CP_TOTAL)
                continue
            flags = model_flags(cl + "|" + da + "|" + date + "|" + tag, da_mul.get((cl, da), 1.0), date, asof)
            out.append([cl, da, date, verdict_for(flags, date, asof), 1, tag, "", ""] + flags)
    return out


def month_range(spec):
    """'2026-05,2026-06' or '2025-10:2026-07' → list of YYYY-MM."""
    out = []
    for part in spec.split(","):
        if not part:
            continue
        if ":" in part:
            a, b = part.split(":")
            y, m = int(a[:4]), int(a[5:7])
            while "%04d-%02d" % (y, m) <= b:
                out.append("%04d-%02d" % (y, m))
                y, m = (y + 1, 1) if m == 12 else (y, m + 1)
        else:
            out.append(part)
    return out


def seed_ratings(rows, months, da_mul):
    """[clinic, da, month, band, remarks, by, at]: a rating for every month
    in which a DA had scoreable cases in their D90 journey (surgery month to
    three months after). The band follows the DA's follow-up level, with
    some month-to-month movement."""
    das = sorted(set((c[0], c[1]) for c in rows if c[4] == 1))
    active = defaultdict(set)
    for c in rows:
        if c[4] != 1:
            continue
        y, m = int(c[2][:4]), int(c[2][5:7])
        for step in range(4):
            active[(c[0], c[1])].add("%04d-%02d" % (y + (m + step - 1) // 12, (m + step - 1) % 12 + 1))
    out = []
    for cl, da in das:
        level = da_mul.get((cl, da), 1.0)
        for m in months:
            if m not in active[(cl, da)]:
                continue
            r = rng("rating|" + cl + "|" + da + "|" + m)
            score = level + (r.random() - 0.5) * 0.24
            band = 1 if score >= 1.12 else 2 if score >= 0.95 else 3
            if band == 3 and score < 0.80 and r.random() < 0.35:   # band 4 is rare
                band = 4
            remark = REMARKS[band][r.randrange(len(REMARKS[band]))]
            y, mo = int(m[:4]), int(m[5:7])
            nxt = datetime.date(y + (mo == 12), (mo % 12) + 1, 2 + r.randrange(4))
            at = datetime.datetime(nxt.year, nxt.month, nxt.day, 2 + r.randrange(8), 5 * r.randrange(12))
            out.append([cl, da, m, band, remark, "Thao", at.strftime("%Y-%m-%dT%H:%M:00Z")])
    return out


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--cte", required=True)
    ap.add_argument("--audit")
    ap.add_argument("--base", default="lasik/index.html")
    ap.add_argument("--write")
    ap.add_argument("--json")
    ap.add_argument("--clinic", default="CTE")
    ap.add_argument("--asof")
    ap.add_argument("--phones", choices=["last4", "full", "none"], default="last4")
    ap.add_argument("--model-from")
    ap.add_argument("--no-model", action="store_true")
    ap.add_argument("--synth", action="append", default=[], metavar="YYYY-MM=N")
    ap.add_argument("--seed-ratings", default="", metavar="YYYY-MM,...")
    a = ap.parse_args()
    if not a.write and not a.json:
        ap.error("give --write or --json")

    txt, m, base = load_base(a.base)
    asof = a.asof or base.get("asof") or ""
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", asof):
        ap.error("--asof YYYY-MM-DD is required (base page has none)")

    rows = carry_over(base, a.clinic)
    da_mul = da_multipliers(rows, asof)
    synth_n = 0
    for spec in a.synth:
        mo, n = spec.split("=")
        new = synth_cases(rows, mo, int(n), da_mul, asof, a.clinic)
        rows += new
        synth_n += len(new)

    cte, info = read_cte(a.cte, a.clinic, a.phones)
    late = [c for c in cte if c["date"] > asof]
    cte = [c for c in cte if c["date"] <= asof]
    audited = apply_audit(cte, read_audit(a.audit)) if a.audit else 0
    modelled = 0
    if not a.no_model:
        modelled = model_cases(cte, da_mul, a.model_from or "0000-00-00", asof)

    for c in sorted(cte, key=lambda c: (c["date"], c["name"])):
        rows.append([c["clinic"], c["da"], c["date"], c["verdict"],
                     0 if c["verdict"] in ("pending", "nozalo") else 1,
                     "", c["name"], c["show"]] + c["flags"])
    rows.sort(key=lambda r: (r[2], r[0], r[1]))

    months = month_range(a.seed_ratings)
    ratings = seed_ratings(rows, months, da_mul) if months else []

    payload = {"v": 2, "asof": asof, "generated": datetime.date.today().isoformat(),
               "touch": TOUCH, "ratings": ratings, "cases": rows}
    out = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))

    if a.json:
        open(a.json, "w", encoding="utf-8").write(out)
    if a.write:
        new = txt[:m.start(2)] + out + txt[m.end(2):]
        open(a.write, "w", encoding="utf-8").write(new)

    by_v = Counter(c["verdict"] for c in cte)
    print("%s: %d cases from %s (%s), zalo status known for %d phones"
          % (a.clinic, len(cte), a.cte, info["sheet"], info["zaloKnown"]))
    print("  verdicts:", dict(by_v), "| audited:", audited, "| modelled:", modelled,
          "| after as-of %s, left out: %d" % (asof, len(late)))
    print("  other clinics: %d carried over + %d made up | ratings embedded: %d | total cases: %d | payload %d bytes"
          % (len(rows) - len(cte) - synth_n, synth_n, len(ratings), len(rows), len(out.encode("utf-8"))))
    if da_mul:
        print("  DA levels:", ", ".join("%s/%s %.2f" % (k[0], k[1], v) for k, v in sorted(da_mul.items())))


if __name__ == "__main__":
    main()
